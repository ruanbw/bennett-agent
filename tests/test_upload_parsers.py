"""解析器：字节 → 文本。

这里不联网、不碰数据库——解析是纯函数（虽然跑在线程里）。重点是**失败方式**：

- 损坏/加密/扫描件的错误必须可读，且能被归类成「不可重试」；
- 编码回退不能把 UTF-8 中文变成乱码却「看起来成功」；
- 归一化必须清掉 NUL（Postgres 的 TEXT 会直接拒绝含 `\\x00` 的写入）。

PDF 的正向用例需要一份真实带文字层的文件；这里只覆盖「损坏 → 明确报错」，
真实 PDF 的端到端验证放在手工验证与后续的 fixture 里（避免把一段二进制塞进仓库）。
"""

from __future__ import annotations

import io

import docx
import pytest

from app.services.uploads.parsers import (
    DocumentParseError,
    extract_text,
    normalize_text,
)
from app.services.uploads.validators import KIND_DOCX, KIND_HTML, KIND_PDF, KIND_TEXT


def _docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for line in paragraphs:
        document.add_paragraph(line)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, value in enumerate(row):
                grid.cell(row_index, column_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_plain_text_uses_utf8_first():
    text = extract_text(KIND_TEXT, "# 标题\n\n退款七天。".encode())

    assert "退款七天" in text


def test_plain_text_falls_back_to_gb18030():
    """GBK 文档不该因为「不是 UTF-8」就变成乱码或失败。"""
    text = extract_text(KIND_TEXT, "退款政策".encode("gb18030"))

    assert "退款政策" in text


def test_html_strips_scripts_and_styles():
    html = (
        "<html><head><style>body{color:red}</style></head><body>"
        "<script>var refund=1</script><h1>退款</h1><p>七天</p></body></html>"
    ).encode()

    text = extract_text(KIND_HTML, html)

    assert "退款" in text and "七天" in text
    # 脚本里的变量名会污染检索（`refund` 误命中），必须被剥掉
    assert "var refund" not in text
    assert "color:red" not in text


def test_docx_extracts_paragraphs_and_tables():
    data = _docx_bytes(["发票规则", "抬头必须是公司全称"], table=[["字段", "要求"], ["税号", "18 位"]])

    text = extract_text(KIND_DOCX, data)

    assert "发票规则" in text
    assert "税号 | 18 位" in text


def test_corrupt_pdf_reports_a_readable_error():
    with pytest.raises(DocumentParseError, match="损坏|格式"):
        extract_text(KIND_PDF, b"%PDF-1.7\nnot really a pdf")


def test_empty_extraction_is_an_explicit_error():
    """扫描版 PDF 抽不出文字时，必须说清原因，而不是入库一篇空文档。"""
    with pytest.raises(DocumentParseError, match="扫描版"):
        extract_text(KIND_TEXT, b"   \n\n  ")


def test_normalize_removes_nul_and_collapses_blank_lines():
    """NUL 会让 Postgres 的 TEXT 写入直接失败（`invalid byte sequence`）。"""
    cleaned = normalize_text("a\x00b\r\n\r\n\r\n\r\nc   \n")

    assert "\x00" not in cleaned
    assert cleaned == "ab\n\n\nc"
